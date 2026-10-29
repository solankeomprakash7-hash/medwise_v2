#!/usr/bin/env python3
"""Checks knowledge/drug_spectrum.csv AWaRe groups against the WHO AWaRe 2023 spreadsheet.

Every filled `aware` value must cite a row ("sheet '<name>', row <n>") in `aware_source`, and that row of
the WHO file must name the same group (or appear on the "Not recommended" sheet).

Usage (needs openpyxl; the spreadsheet is kept on the private `sources` branch):
    python3 scripts/verify-aware.py /path/to/WHO-MHP-HPS-EML-2023.04-eng.xlsx
"""
import csv
import re
import sys

import openpyxl

workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
sheets = {name: list(workbook[name].iter_rows(values_only=True)) for name in workbook.sheetnames}

failures = 0
rows = list(csv.DictReader(open("knowledge/drug_spectrum.csv", encoding="utf8")))
for row in rows:
    drug, group, source = row["drug"], row["aware"], row["aware_source"]
    if not group:
        print(f"empty {drug}: {source}")
        continue
    cited = re.search(r"sheet '([^']+)', row (\d+)", source)
    if not cited:
        print(f"FAIL  {drug}: no sheet and row cited")
        failures += 1
        continue
    sheet, number = cited.group(1), int(cited.group(2))
    who_row = sheets[sheet][number - 1]
    who_group = "Not recommended" if sheet == "Not recommended" else str(who_row[3]).strip()
    if who_group != group:
        print(f"FAIL  {drug}: table says {group}, WHO {sheet} row {number} ({who_row[0]}) says {who_group}")
        failures += 1
    else:
        print(f"ok    {drug}: {group} ({who_row[0]})")

filled = sum(1 for r in rows if r["aware"])
print(f"\n{filled - failures} of {filled} AWaRe groups match the WHO spreadsheet.")
sys.exit(1 if failures else 0)
